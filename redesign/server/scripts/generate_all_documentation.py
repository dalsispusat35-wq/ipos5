# -*- coding: utf-8 -*-
"""
Script Generator & Compiler Dokumentasi Lengkap Halaman IPOS5 PT Pos Indonesia
Menghasilkan 16 file Markdown (.md) dan 16 file Microsoft Word (.docx) serta 1 file kompilasi master DOCX.
Masing-masing halaman dilengkapi Wireframe ASCII Tata Letak UI dan Penjelasan Rinci Setiap Kotak/Komponen.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = r"c:\Users\Asus\Documents\POSIND\IPOS5\ipos5\dokumentasi_halaman_ipos5"
os.makedirs(DOCS_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_docx_for_page(page_dict, file_path):
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.0)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(page_dict["title"])
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    if "subtitle" in page_dict and page_dict["subtitle"]:
        p_sub = doc.add_paragraph()
        r_sub = p_sub.add_run(page_dict["subtitle"])
        r_sub.font.size = Pt(11.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = RGBColor(0xCC, 0x33, 0x00)
    
    if "metadata" in page_dict and page_dict["metadata"]:
        h_meta = doc.add_heading(level=2)
        r_hm = h_meta.add_run("📌 METADATA & INTEGRASI SISTEM")
        r_hm.font.size = Pt(10.5)
        r_hm.font.bold = True
        r_hm.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        
        t_meta = doc.add_table(rows=len(page_dict["metadata"]) + 1, cols=2)
        t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths_m = [Inches(2.2), Inches(4.6)]
        
        hdr_m = t_meta.rows[0].cells
        hdr_m[0].text = "Parameter"
        hdr_m[1].text = "Nilai / Spesifikasi Sistem"
        for idx in range(2):
            hdr_m[idx].width = widths_m[idx]
            set_cell_background(hdr_m[idx], "003366")
            set_cell_margins(hdr_m[idx])
            p = hdr_m[idx].paragraphs[0]
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9.0)
                
        for row_idx, (k, v) in enumerate(page_dict["metadata"]):
            row_cells = t_meta.rows[row_idx + 1].cells
            row_cells[0].width = widths_m[0]
            row_cells[1].width = widths_m[1]
            row_cells[0].text = k
            row_cells[1].text = v
            bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
            for c_idx in range(2):
                set_cell_background(row_cells[c_idx], bg)
                set_cell_margins(row_cells[c_idx])
                p = row_cells[c_idx].paragraphs[0]
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    if c_idx == 0:
                        r.font.bold = True
        doc.add_paragraph()
    
    for sec in page_dict.get("sections", []):
        h = doc.add_heading(level=2)
        r_h = h.add_run(sec["heading"])
        r_h.font.size = Pt(11.5)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        
        if "content" in sec and sec["content"]:
            p_c = doc.add_paragraph()
            r_c = p_c.add_run(sec["content"])
            r_c.font.size = Pt(9.5)
            
        if "wireframe" in sec and sec["wireframe"]:
            p_wf_title = doc.add_paragraph()
            r_wft = p_wf_title.add_run("📐 Diagram Visual Tata Letak Halaman (Wireframe UI Layout):")
            r_wft.font.bold = True
            r_wft.font.size = Pt(9.5)
            r_wft.font.color.rgb = RGBColor(0x00, 0x44, 0x88)
            
            p_wf = doc.add_paragraph()
            r_wf = p_wf.add_run(sec["wireframe"])
            r_wf.font.name = 'Consolas'
            r_wf.font.size = Pt(8.0)
            p_wf.paragraph_format.left_indent = Inches(0.1)
            p_wf.paragraph_format.space_before = Pt(2)
            p_wf.paragraph_format.space_after = Pt(4)
            
        if "components_table" in sec and sec["components_table"]:
            p_ct_title = doc.add_paragraph()
            r_ctt = p_ct_title.add_run("📋 Rincian & Keterangan Setiap Kotak / Komponen Tata Letak:")
            r_ctt.font.bold = True
            r_ctt.font.size = Pt(10.0)
            r_ctt.font.color.rgb = RGBColor(0xCC, 0x33, 0x00)
            
            headers = ["Bagian / Kotak UI", "Tipe Visual & Posisi", "Fungsi & Kegunaan", "Sumber Data & Kalkulasi", "Interaktivitas & Aksi"]
            widths_c = [Inches(1.4), Inches(1.2), Inches(1.6), Inches(1.4), Inches(1.2)]
            t_comp = doc.add_table(rows=len(sec["components_table"]) + 1, cols=5)
            t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            hdr_cells = t_comp.rows[0].cells
            for idx, th in enumerate(headers):
                hdr_cells[idx].width = widths_c[idx]
                hdr_cells[idx].text = th
                set_cell_background(hdr_cells[idx], "003366")
                set_cell_margins(hdr_cells[idx])
                p = hdr_cells[idx].paragraphs[0]
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(8.0)
                    
            for r_idx, comp_row in enumerate(sec["components_table"]):
                row_cells = t_comp.rows[r_idx + 1].cells
                bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx in range(5):
                    row_cells[c_idx].width = widths_c[c_idx]
                    row_cells[c_idx].text = comp_row[c_idx]
                    set_cell_background(row_cells[c_idx], bg)
                    set_cell_margins(row_cells[c_idx])
                    p = row_cells[c_idx].paragraphs[0]
                    for r in p.runs:
                        r.font.size = Pt(7.5)
                        if c_idx == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            doc.add_paragraph()
            
        if "items" in sec and sec["items"]:
            for item in sec["items"]:
                p_it = doc.add_paragraph(style='List Bullet')
                r_it = p_it.add_run(item)
                r_it.font.size = Pt(9.0)
                
    doc.save(file_path)
    print(f"[OK DOCX] {os.path.basename(file_path)}")

def write_markdown_file(page_dict, file_path):
    lines = []
    lines.append(f"# {page_dict['title']}")
    if "subtitle" in page_dict and page_dict["subtitle"]:
        lines.append(f"## {page_dict['subtitle']}")
    lines.append("")
    
    if "metadata" in page_dict and page_dict["metadata"]:
        lines.append("### 📌 METADATA & INTEGRASI SISTEM")
        lines.append("| Parameter | Nilai / Spesifikasi |")
        lines.append("| :--- | :--- |")
        for k, v in page_dict["metadata"]:
            lines.append(f"| **{k}** | {v} |")
        lines.append("")
        lines.append("---")
        lines.append("")
        
    for sec in page_dict.get("sections", []):
        lines.append(f"## {sec['heading']}")
        lines.append("")
        if "content" in sec and sec["content"]:
            lines.append(sec["content"])
            lines.append("")
            
        if "wireframe" in sec and sec["wireframe"]:
            lines.append("### 📐 Diagram Visual Tata Letak Halaman (Wireframe UI Layout)")
            lines.append("```text")
            lines.append(sec["wireframe"])
            lines.append("```")
            lines.append("")
            
        if "components_table" in sec and sec["components_table"]:
            lines.append("### 📋 Rincian & Keterangan Setiap Kotak / Komponen Tata Letak")
            lines.append("| Bagian / Kotak UI | Tipe Visual & Posisi | Fungsi & Kegunaan | Sumber Data & Kalkulasi | Interaktivitas & Aksi |")
            lines.append("| :--- | :--- | :--- | :--- | :--- |")
            for row in sec["components_table"]:
                c0 = row[0].replace("|", "\\|")
                c1 = row[1].replace("|", "\\|")
                c2 = row[2].replace("|", "\\|")
                c3 = row[3].replace("|", "\\|")
                c4 = row[4].replace("|", "\\|")
                lines.append(f"| **{c0}** | {c1} | {c2} | {c3} | {c4} |")
            lines.append("")
            
        if "items" in sec and sec["items"]:
            for it in sec["items"]:
                lines.append(f"- {it}")
            lines.append("")
            
    with open(file_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"[OK MD] {os.path.basename(file_path)}")

print("Core generator functions ready.")
